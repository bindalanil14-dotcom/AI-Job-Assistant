"""
RAG pipeline for the AI Job Application Assistant.

Responsibilities:
- Extract plain text from uploaded documents (.txt, .pdf, .docx).
- Chunk text into overlapping windows.
- Embed text with OpenAI text-embedding-3-small.
- Compute cosine similarity and retrieve the most relevant chunks.
- Provide a reusable semantic-similarity helper used by the skill matcher so that
  matching goes BEYOND simple keyword overlap (e.g. "ML" ~ "Machine Learning").
"""

import io
import os
from openai import OpenAI

EMBED_MODEL = "text-embedding-3-small"


# ----------------------------------------------------------------------------
# TEXT EXTRACTION (.txt / .pdf / .docx)
# ----------------------------------------------------------------------------
def extract_text_from_upload(file_name, file_bytes):
    """
    Extract plain text from an uploaded document given its name and raw bytes.
    Supports .txt, .pdf and .docx. Raises ValueError for unsupported types.
    """
    ext = os.path.splitext(file_name)[1].lower()
    if ext == ".txt":
        return _extract_txt(file_bytes)
    if ext == ".pdf":
        return _extract_pdf(file_bytes, file_name)
    if ext == ".docx":
        return _extract_docx(file_bytes)
    raise ValueError(f"Unsupported file type '{ext}'. Use .txt, .pdf or .docx.")


def _extract_txt(file_bytes):
    for enc in ("utf-8", "utf-16", "latin-1"):
        try:
            text = file_bytes.decode(enc)
            if text.strip():
                return text
        except Exception:
            continue
    return file_bytes.decode("utf-8", errors="ignore")


def _extract_pdf(file_bytes, file_name):
    text = ""
    # Try pdfplumber first
    try:
        import pdfplumber
        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            for page in pdf.pages:
                val = page.extract_text()
                if val:
                    text += val + "\n"
        if text.strip():
            return text
    except Exception:
        pass

    # Fallback to PyPDF2
    try:
        import PyPDF2
        reader = PyPDF2.PdfReader(io.BytesIO(file_bytes))
        for page in reader.pages:
            val = page.extract_text()
            if val:
                text += val + "\n"
        if text.strip():
            return text
    except Exception as e:
        raise RuntimeError(
            f"Failed to extract text from PDF '{file_name}': {str(e)}"
        )

    if not text.strip():
        raise ValueError(f"No text could be extracted from '{file_name}'.")
    return text


def _extract_docx(file_bytes):
    try:
        import docx  # python-docx
    except Exception as e:
        raise RuntimeError(
            "python-docx is required to read .docx files. Install with "
            "`pip install python-docx`. Original error: " + str(e)
        )
    document = docx.Document(io.BytesIO(file_bytes))
    parts = [p.text for p in document.paragraphs if p.text]
    # Include table cell text too (resumes often use tables for layout)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text:
                    parts.append(cell.text)
    text = "\n".join(parts)
    if not text.strip():
        raise ValueError("No text could be extracted from the .docx file.")
    return text


# ----------------------------------------------------------------------------
# CHUNKING
# ----------------------------------------------------------------------------
def chunk_text(text, chunk_size=600, overlap=100):
    """Chunk text into overlapping character windows."""
    chunks = []
    if not text:
        return chunks
    text = " ".join(text.split())  # normalise whitespace for cleaner chunks
    step = chunk_size - overlap
    if step <= 0:
        step = chunk_size
    for i in range(0, len(text), step):
        chunk = text[i:i + chunk_size].strip()
        if chunk:
            chunks.append(chunk)
    return chunks


# ----------------------------------------------------------------------------
# EMBEDDINGS + SIMILARITY
# ----------------------------------------------------------------------------
def cosine_similarity(v1, v2):
    dot = sum(x * y for x, y in zip(v1, v2))
    n1 = sum(x * x for x in v1) ** 0.5
    n2 = sum(y * y for y in v2) ** 0.5
    if n1 == 0 or n2 == 0:
        return 0.0
    return dot / (n1 * n2)


def get_embeddings_batch(texts, api_key, batch_size=100):
    """Embed a list of strings via the modular provider. Aligned to `texts`."""
    if not texts:
        return []
    from llm_provider import get_provider
    return get_provider(api_key).embed(texts, batch_size=batch_size)


def retrieve_relevant_chunks(query, chunks, chunk_embeddings, api_key, top_k=5):
    """Return the top_k chunks most semantically similar to the query."""
    if not chunks or not chunk_embeddings:
        return []
    from llm_provider import get_provider
    q_emb = get_provider(api_key).embed([query])[0]
    scored = []
    for chunk, emb in zip(chunks, chunk_embeddings):
        scored.append((cosine_similarity(q_emb, emb), chunk))
    scored.sort(key=lambda x: x[0], reverse=True)
    return [c for _, c in scored[:top_k]]


# ----------------------------------------------------------------------------
# COLOR EXTRACTION FROM A REFERENCE RESUME (real colors, not inferred)
# ----------------------------------------------------------------------------
def extract_colors_from_upload(file_name, file_bytes):
    """
    Parse the dominant accent colour actually used in a reference resume so the
    generated resume can match its visual template.

    PDF  -> pdfplumber char/rect fill colours.
    DOCX -> python-docx run font colours.
    Returns {"accent","heading","name"} hex strings, or {} if nothing usable.
    """
    ext = os.path.splitext(file_name)[1].lower()
    try:
        if ext == ".pdf":
            return _colors_from_pdf(file_bytes)
        if ext == ".docx":
            return _colors_from_docx(file_bytes)
    except Exception:
        return {}
    return {}


def _to_rgb(color):
    """Normalise a pdfplumber/python-docx colour value to an (r,g,b) 0-255 tuple."""
    if color is None:
        return None
    if isinstance(color, (int, float)):
        vals = [color, color, color]
    else:
        try:
            seq = list(color)
        except TypeError:
            return None
        if len(seq) == 1:
            vals = [seq[0]] * 3
        elif len(seq) == 3:
            vals = seq
        elif len(seq) == 4:  # CMYK
            c, m, y, k = seq
            vals = [(1 - c) * (1 - k), (1 - m) * (1 - k), (1 - y) * (1 - k)]
        else:
            return None
    try:
        if max(vals) > 1:  # already 0-255
            vals = [v / 255.0 for v in vals]
    except TypeError:
        return None
    return tuple(max(0, min(255, int(round(v * 255)))) for v in vals)


def _is_meaningful(rgb):
    """Skip near-black, near-white and greys — keep real brand/accent colours."""
    mx, mn = max(rgb), min(rgb)
    if mx < 60:        # near black (body text)
        return False
    if mn > 200:       # near white (background)
        return False
    if mx - mn < 28:   # greyscale
        return False
    return True


def _dominant_color(buckets):
    """buckets: key -> [count, sum_r, sum_g, sum_b]. Return averaged hex of the top bucket."""
    if not buckets:
        return {}
    key = max(buckets, key=lambda k: buckets[k][0])
    cnt, sr, sg, sb = buckets[key]
    r, g, b = int(sr / cnt), int(sg / cnt), int(sb / cnt)
    hexv = "#%02x%02x%02x" % (r, g, b)
    return {"accent": hexv, "heading": hexv, "name": hexv}


def _add(buckets, rgb, weight=1):
    key = tuple((v // 16) for v in rgb)  # quantise to group near-identical shades
    b = buckets.setdefault(key, [0, 0, 0, 0])
    b[0] += weight
    b[1] += rgb[0] * weight
    b[2] += rgb[1] * weight
    b[3] += rgb[2] * weight


def _colors_from_pdf(file_bytes):
    import pdfplumber
    buckets = {}
    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        for page in pdf.pages[:4]:
            for ch in page.chars:
                rgb = _to_rgb(ch.get("non_stroking_color"))
                if rgb and _is_meaningful(rgb):
                    _add(buckets, rgb, 1)
            # Coloured banners/sidebars in templates are a strong signal — weight them.
            for rect in page.rects:
                rgb = _to_rgb(rect.get("non_stroking_color"))
                if rgb and _is_meaningful(rgb):
                    _add(buckets, rgb, 6)
    return _dominant_color(buckets)


def _colors_from_docx(file_bytes):
    import docx
    document = docx.Document(io.BytesIO(file_bytes))
    buckets = {}

    def harvest(paragraphs, weight):
        for p in paragraphs:
            for run in p.runs:
                try:
                    val = run.font.color.rgb
                except Exception:
                    val = None
                if val is None:
                    continue
                hx = str(val)
                if len(hx) != 6:
                    continue
                try:
                    rgb = tuple(int(hx[i:i + 2], 16) for i in (0, 2, 4))
                except ValueError:
                    continue
                if _is_meaningful(rgb):
                    _add(buckets, rgb, weight)

    harvest(document.paragraphs, 1)
    # Heading-styled paragraphs often carry the template's accent colour.
    for p in document.paragraphs:
        if p.style and p.style.name and "Head" in p.style.name:
            harvest([p], 5)
    return _dominant_color(buckets)


def semantic_skill_signal(jd_skills, profile_skills, api_key, threshold=0.55):
    """
    For each JD skill, find the most semantically similar profile skill (by embedding
    cosine similarity) and report it with that profile skill's level. This is the
    'beyond keyword matching' signal handed to the Skill Matcher agent.

    `profile_skills` is a list of {"name", "level"} dicts.
    Returns a list of dicts:
      {"jd_skill", "best_profile_skill", "similarity", "level", "matched"}
    """
    if not jd_skills:
        return []
    profile_names = [s.get("name", "") for s in profile_skills if s.get("name")]
    level_by_name = {s["name"]: s.get("level", "moderate") for s in profile_skills if s.get("name")}

    if not profile_names:
        return [
            {"jd_skill": s, "best_profile_skill": None, "similarity": 0.0,
             "level": None, "matched": False}
            for s in jd_skills
        ]

    jd_embs = get_embeddings_batch(jd_skills, api_key)
    prof_embs = get_embeddings_batch(profile_names, api_key)

    signal = []
    for jd_skill, jd_emb in zip(jd_skills, jd_embs):
        best_sim, best_name = 0.0, None
        for name, p_emb in zip(profile_names, prof_embs):
            sim = cosine_similarity(jd_emb, p_emb)
            if sim > best_sim:
                best_sim, best_name = sim, name
        matched = best_sim >= threshold
        signal.append({
            "jd_skill": jd_skill,
            "best_profile_skill": best_name if matched else None,
            "similarity": round(best_sim, 3),
            "level": level_by_name.get(best_name) if matched else None,
            "matched": matched,
        })
    return signal
